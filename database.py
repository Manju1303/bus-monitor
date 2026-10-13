"""
database.py - SQLite Database Logger for Bus Gate Monitoring System
--------------------------------------------------------------------
Thread-safe SQLite backend replacing Excel for gate activity storage.
10x faster writes, no corruption risk, supports concurrent access.
Excel export is available on demand via export_to_excel().
"""

import os
import threading
import sqlite3
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill


class GateDatabase:
    """
    SQLite-backed gate activity logger.
    Drop-in replacement for ExcelLogger with the same public interface.
    """

    COLUMNS = [
        "Log ID", "Vehicle ID", "Registration Plate", "Direction",
        "Entry Time", "Exit Time", "Duration (Minutes)", "Date", "Status", "Snapshot"
    ]

    def __init__(self, db_path: str = "bus_gate.db"):
        self.db_path = db_path
        self._lock = threading.Lock()
        self._init_db()

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _connect(self):
        """Return a new SQLite connection with row_factory set."""
        conn = sqlite3.connect(self.db_path, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self):
        """Create the gate_logs table if it does not already exist."""
        with self._lock:
            conn = self._connect()
            conn.execute("""
                CREATE TABLE IF NOT EXISTS gate_logs (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    vehicle_id  TEXT,
                    plate       TEXT,
                    direction   TEXT,
                    entry_time  TEXT,
                    exit_time   TEXT,
                    duration    REAL,
                    date        TEXT,
                    status      TEXT,
                    snapshot    TEXT
                )
            """)
            conn.commit()
            conn.close()

    def _ensure_file_exists(self):
        """Compatibility shim — SQLite DB initialises itself automatically."""
        self._init_db()

    # ------------------------------------------------------------------
    # Write operations
    # ------------------------------------------------------------------

    def log_vehicle(self, vehicle_id: int, plate_no: str,
                    direction: str, snapshot: str = None) -> dict:
        """
        Log an Entry or Exit gate event.
        - For Exit: finds the matching ON CAMPUS row and updates it.
        - For Entry (or unmatched Exit): inserts a new row.

        :param vehicle_id: Tracker-assigned numeric vehicle ID
        :param plate_no:   Recognised plate string (e.g. 'TN38AB1234')
        :param direction:  'Entry' or 'Exit'
        :param snapshot:   Optional filename of JPEG snapshot saved to snapshots/
        :return:           Summary dict of the logged event
        """
        now = datetime.now()
        timestamp_str = now.strftime("%Y-%m-%d %H:%M:%S")
        date_str = now.strftime("%Y-%m-%d")
        plate_clean = plate_no.upper() if plate_no else "UNKNOWN"

        with self._lock:
            conn = self._connect()
            try:
                updated = False

                if direction == "Exit":
                    row = conn.execute(
                        """SELECT id, entry_time FROM gate_logs
                           WHERE plate = ? AND status = 'ON CAMPUS'
                           ORDER BY id DESC LIMIT 1""",
                        (plate_clean,)
                    ).fetchone()

                    if row:
                        try:
                            entry_dt = datetime.strptime(row["entry_time"], "%Y-%m-%d %H:%M:%S")
                            duration_mins = round((now - entry_dt).total_seconds() / 60.0, 1)
                        except Exception:
                            duration_mins = 0.0

                        conn.execute(
                            """UPDATE gate_logs
                               SET exit_time = ?, duration = ?, status = 'EXITED', snapshot = COALESCE(?, snapshot)
                               WHERE id = ?""",
                            (timestamp_str, duration_mins, snapshot, row["id"])
                        )
                        conn.commit()
                        updated = True

                if not updated:
                    status  = "ON CAMPUS" if direction == "Entry" else "EXITED"
                    entry_t = timestamp_str if direction == "Entry" else None
                    exit_t  = timestamp_str if direction == "Exit"  else None

                    conn.execute(
                        """INSERT INTO gate_logs
                           (vehicle_id, plate, direction, entry_time, exit_time,
                            duration, date, status, snapshot)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (f"BUS-{vehicle_id}", plate_clean, direction,
                         entry_t, exit_t, None, date_str, status, snapshot)
                    )
                    conn.commit()

            finally:
                conn.close()

        print(f"[DB LOG] {direction} — {plate_clean} at {timestamp_str}")
        return {
            "vehicle_id": vehicle_id,
            "plate":      plate_clean,
            "direction":  direction,
            "timestamp":  timestamp_str
        }

    # ------------------------------------------------------------------
    # Read operations
    # ------------------------------------------------------------------

    def get_summary_stats(self) -> dict:
        """Return KPI counts: total records, currently on campus, exited today."""
        today = datetime.now().strftime("%Y-%m-%d")
        with self._lock:
            conn = self._connect()
            try:
                total       = conn.execute("SELECT COUNT(*) FROM gate_logs").fetchone()[0]
                on_campus   = conn.execute(
                    "SELECT COUNT(*) FROM gate_logs WHERE status = 'ON CAMPUS'"
                ).fetchone()[0]
                exited_today = conn.execute(
                    "SELECT COUNT(*) FROM gate_logs WHERE status = 'EXITED' AND date = ?",
                    (today,)
                ).fetchone()[0]
            finally:
                conn.close()

        return {
            "total_records": total,
            "on_campus":     on_campus,
            "exited_today":  exited_today,
            "model_status":  "Active (YOLOv8 + EasyOCR)"
        }

    def get_records_list(self, search_query: str = "", status_filter: str = "ALL") -> list:
        """
        Return gate log records as a list of dicts, with optional filtering.

        :param search_query:  Text to match against plate, vehicle_id, or direction
        :param status_filter: 'ON CAMPUS', 'EXITED', or 'ALL'
        """
        conditions = []
        params = []

        if status_filter and status_filter.upper() != "ALL":
            conditions.append("status = ?")
            params.append(status_filter.upper())

        if search_query:
            q = f"%{search_query.strip()}%"
            conditions.append("(plate LIKE ? OR vehicle_id LIKE ? OR direction LIKE ?)")
            params.extend([q, q, q])

        sql = "SELECT * FROM gate_logs"
        if conditions:
            sql += " WHERE " + " AND ".join(conditions)
        sql += " ORDER BY id DESC"

        with self._lock:
            conn = self._connect()
            try:
                rows = conn.execute(sql, params).fetchall()
            finally:
                conn.close()

        records = []
        for row in rows:
            records.append({
                "Log ID":              row["id"],
                "Vehicle ID":          row["vehicle_id"] or "-",
                "Registration Plate":  row["plate"] or "-",
                "Direction":           row["direction"] or "-",
                "Entry Time":          row["entry_time"] or "-",
                "Exit Time":           row["exit_time"] or "-",
                "Duration (Minutes)":  row["duration"] if row["duration"] is not None else "-",
                "Date":                row["date"] or "-",
                "Status":              row["status"] or "-",
                "Snapshot":            row["snapshot"] or ""
            })
        return records

    # ------------------------------------------------------------------
    # Excel export
    # ------------------------------------------------------------------

    def export_to_excel(self, output_path: str = "bus_gate_log.xlsx") -> str:
        """
        Export all gate logs from SQLite to a styled Excel file.

        :param output_path: Destination .xlsx file path
        :return:            The output path (for chaining)
        """
        records = self.get_records_list()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "JKKM Gate Activity Logs"

        export_cols = [
            "Log ID", "Vehicle ID", "Registration Plate", "Direction",
            "Entry Time", "Exit Time", "Duration (Minutes)", "Date", "Status"
        ]
        ws.append(export_cols)

        # Header styling
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(export_cols) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill   = header_fill
            cell.font   = header_font
            cell.alignment = align_center

        widths = [10, 15, 22, 14, 20, 20, 20, 15, 15]
        for i, letter in enumerate("ABCDEFGHI"):
            ws.column_dimensions[letter].width = widths[i]

        # Data rows
        for rec in records:
            row_data = [rec.get(c, "-") for c in export_cols]
            ws.append(row_data)
            for col_num in range(1, len(row_data) + 1):
                ws.cell(row=ws.max_row, column=col_num).alignment = align_center

        wb.save(output_path)
        print(f"[EXPORT] Excel report saved: {output_path} ({len(records)} records)")
        return output_path
